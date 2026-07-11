#!/usr/bin/env python3
from __future__ import annotations

import getpass
import csv
import json
import math
import os
import re
import selectors
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Protocol

import requests
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


BRANCH_SHEET = "branch"
LATITUDE_COLUMN = "latitude"
LONGITUDE_COLUMN = "longitude"
FULL_ADDRESS_COLUMN = "full_address"
PARSED_ADDRESS_COLUMNS = [
    "barangay2",
    "city_municipality2",
    "province2",
    "region2",
    "zip_code2",
    "country2",
    "island_group2",
]
ADDRESS_FIELDS = [
    "barangay",
    "city_municipality",
    "province",
    "region",
    "zip_code",
    "country",
    "island_group",
]
DEFAULT_GEMINI_MODEL = "gemini-3.1-flash-lite"
NULL_LIKE_VALUES = {"", "none", "null", "nan", "n/a", "na", "-", "--"}
MODELS_DIR = Path(__file__).resolve().parent / "models"
MODEL_INDEX_FILENAME = "models_index.json"
LOCAL_MODEL_TOO_LARGE_BYTES = 4 * 1024 * 1024 * 1024
LOCAL_WORKER_TIMEOUT_SECONDS = 180
LOCAL_DEFAULT_THREADS = min(8, os.cpu_count() or 4)


class BranchCleanerError(Exception):
    """Raised for user-actionable workbook/input problems."""


class Geocoder(Protocol):
    def geocode(self, address: str) -> tuple[float, float] | None:
        """Return latitude and longitude for an address, or None if no result."""


class AddressParser(Protocol):
    def parse_address(self, address: str) -> dict[str, str | None]:
        """Return normalized address fields for an address."""


@dataclass
class CoordinateCleanSummary:
    rows_checked: int = 0
    rows_needing_coordinates: int = 0
    geocoded_rows: int = 0
    skipped_no_address_rows: list[int] = field(default_factory=list)
    no_result_rows: list[int] = field(default_factory=list)
    failed_rows: dict[int, str] = field(default_factory=dict)


@dataclass
class ParseAddressSummary:
    rows_checked: int = 0
    parsed_rows: int = 0
    skipped_no_address_rows: list[int] = field(default_factory=list)
    failed_rows: dict[int, str] = field(default_factory=dict)


@dataclass(frozen=True)
class CoordinateColumns:
    latitude: int
    longitude: int
    address: int


@dataclass(frozen=True)
class HeaderColumn:
    label: str
    index: int


@dataclass(frozen=True)
class HuggingFaceGgufFile:
    repo_id: str
    filename: str
    size_bytes: int | None = None


@dataclass(frozen=True)
class LocalGgufModel:
    repo_id: str
    filename: str
    path: Path
    size_bytes: int | None = None


@dataclass
class LoadedFile:
    path: Path
    kind: str
    workbook: object
    sheet: Worksheet


class GoogleGeocoder:
    def __init__(self, api_key: str, session: requests.Session | None = None):
        self.api_key = api_key
        self.session = session or requests.Session()

    def geocode(self, address: str) -> tuple[float, float] | None:
        response = self.session.get(
            "https://maps.googleapis.com/maps/api/geocode/json",
            params={"address": address, "key": self.api_key},
            timeout=30,
        )
        response.raise_for_status()
        payload = response.json()
        if payload.get("status") != "OK" or not payload.get("results"):
            return None
        location = payload["results"][0].get("geometry", {}).get("location", {})
        lat = parse_coordinate(location.get("lat"), "latitude")
        lng = parse_coordinate(location.get("lng"), "longitude")
        if lat is None or lng is None:
            return None
        return lat, lng


class GeminiApiAddressParser:
    def __init__(self, api_key: str, model: str = DEFAULT_GEMINI_MODEL, session: requests.Session | None = None):
        self.api_key = api_key
        self.model = model
        self.session = session or requests.Session()

    def parse_address(self, address: str) -> dict[str, str | None]:
        payload = {
            "contents": [{"parts": [{"text": build_address_prompt(address)}]}],
            "generationConfig": {
                "responseMimeType": "application/json",
                "responseSchema": address_json_schema(),
            },
        }
        response = self.session.post(
            f"https://generativelanguage.googleapis.com/v1beta/models/{self.model}:generateContent",
            params={"key": self.api_key},
            headers={"Content-Type": "application/json"},
            json=payload,
            timeout=60,
        )
        response.raise_for_status()
        response_payload = response.json()
        raw = extract_gemini_json_payload(response_payload)
        return normalize_address_parse_result(raw)


class LlamaCppAddressParser:
    def __init__(
        self,
        model_path: Path,
        llama_factory: Callable[..., object] | None = None,
        *,
        n_ctx: int = 1536,
        n_gpu_layers: int = -1,
        n_batch: int = 512,
        n_ubatch: int = 128,
        n_threads: int | None = None,
        n_threads_batch: int | None = None,
        flash_attn: bool = False,
    ):
        self.model_path = Path(model_path)
        if not self.model_path.exists():
            raise BranchCleanerError(f"Local model file does not exist: {self.model_path}")
        self.llama_factory = llama_factory or load_llama_cpp_factory()
        self.model = self.llama_factory(
            model_path=str(self.model_path),
            n_ctx=n_ctx,
            n_gpu_layers=n_gpu_layers,
            n_batch=n_batch,
            n_ubatch=n_ubatch,
            n_threads=n_threads,
            n_threads_batch=n_threads_batch,
            flash_attn=flash_attn,
            verbose=False,
        )

    def parse_address(self, address: str) -> dict[str, str | None]:
        first_text = run_llama_address_chat(self.model, address)
        try:
            return normalize_address_parse_result(first_text)
        except Exception:
            repair_prompt = build_local_json_repair_prompt(address, first_text)
            repair_text = run_llama_repair_chat(self.model, repair_prompt)
            return normalize_address_parse_result(repair_text)


class LlamaCppWorkerAddressParser:
    """Run llama.cpp in a child process so native crashes do not kill the cleaner."""

    def __init__(
        self,
        model_path: Path,
        *,
        n_ctx: int = 1536,
        n_gpu_layers: int = -1,
        n_batch: int = 512,
        n_ubatch: int = 128,
        n_threads: int | None = None,
        n_threads_batch: int | None = None,
        flash_attn: bool = False,
        timeout_seconds: int = LOCAL_WORKER_TIMEOUT_SECONDS,
        python_executable: str | None = None,
        persistent: bool = True,
    ):
        self.model_path = Path(model_path)
        if not self.model_path.exists():
            raise BranchCleanerError(f"Local model file does not exist: {self.model_path}")
        self.n_ctx = n_ctx
        self.n_gpu_layers = n_gpu_layers
        self.n_batch = n_batch
        self.n_ubatch = n_ubatch
        self.n_threads = n_threads
        self.n_threads_batch = n_threads_batch
        self.flash_attn = flash_attn
        self.timeout_seconds = timeout_seconds
        self.python_executable = python_executable or sys.executable
        self.persistent = persistent
        self.process: subprocess.Popen | None = None
        self.stderr_handle = None
        self.stderr_path: Path | None = None

    def parse_address(self, address: str) -> dict[str, str | None]:
        if self.persistent:
            return self._parse_address_persistent(address)
        return self._parse_address_one_shot(address)

    def _parse_address_one_shot(self, address: str) -> dict[str, str | None]:
        payload = {
            "model_path": str(self.model_path),
            "address": address,
            "n_ctx": self.n_ctx,
            "n_gpu_layers": self.n_gpu_layers,
            "n_batch": self.n_batch,
            "n_ubatch": self.n_ubatch,
            "n_threads": self.n_threads,
            "n_threads_batch": self.n_threads_batch,
            "flash_attn": self.flash_attn,
        }
        try:
            completed = subprocess.run(
                [self.python_executable, str(Path(__file__).resolve()), "--local-parse-worker"],
                input=json.dumps(payload),
                text=True,
                capture_output=True,
                timeout=self.timeout_seconds,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise BranchCleanerError(
                f"Local model worker timed out after {self.timeout_seconds}s. "
                "Try a smaller GGUF model or Gemini mode."
            ) from exc

        if completed.returncode != 0:
            stderr_tail = completed.stderr.strip().splitlines()[-8:]
            detail = "\n".join(stderr_tail) or completed.stdout.strip() or "no worker output"
            raise BranchCleanerError(
                f"Local model worker failed with exit code {completed.returncode}. "
                "This often means the selected GGUF is incompatible with the installed llama-cpp-python build, "
                "or Metal/GPU execution crashed. Try CPU mode, a smaller Q4 GGUF, or upgrade llama-cpp-python.\n"
                f"Worker detail:\n{detail}"
            )

        try:
            result = json.loads(completed.stdout)
        except json.JSONDecodeError as exc:
            raise BranchCleanerError(
                "Local model worker returned non-JSON output.\n"
                f"stdout: {completed.stdout[-1000:]}\n"
                f"stderr: {completed.stderr[-1000:]}"
            ) from exc
        if not result.get("ok"):
            raise BranchCleanerError(str(result.get("error") or "Local model worker failed."))
        return normalize_address_parse_result(result.get("parsed"))

    def _parse_address_persistent(self, address: str) -> dict[str, str | None]:
        self._ensure_process()
        assert self.process is not None
        assert self.process.stdin is not None
        request = {"address": address}
        try:
            self.process.stdin.write(json.dumps(request, ensure_ascii=False) + "\n")
            self.process.stdin.flush()
        except BrokenPipeError as exc:
            raise BranchCleanerError(self._worker_failure_message("Local model worker stopped before receiving a row.")) from exc
        response = self._read_worker_json("local model response")
        if not response.get("ok"):
            raise BranchCleanerError(str(response.get("error") or "Local model worker failed."))
        return normalize_address_parse_result(response.get("parsed"))

    def _ensure_process(self) -> None:
        if self.process is not None and self.process.poll() is None:
            return
        self.close()
        stderr_file = tempfile.NamedTemporaryFile(
            mode="w+",
            encoding="utf-8",
            prefix="branch_cleaner_llama_",
            suffix=".log",
            delete=False,
        )
        self.stderr_handle = stderr_file
        self.stderr_path = Path(stderr_file.name)
        self.process = subprocess.Popen(
            [self.python_executable, str(Path(__file__).resolve()), "--local-parse-server"],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=stderr_file,
            text=True,
            bufsize=1,
        )
        assert self.process.stdin is not None
        init_payload = {
            "model_path": str(self.model_path),
            "n_ctx": self.n_ctx,
            "n_gpu_layers": self.n_gpu_layers,
            "n_batch": self.n_batch,
            "n_ubatch": self.n_ubatch,
            "n_threads": self.n_threads,
            "n_threads_batch": self.n_threads_batch,
            "flash_attn": self.flash_attn,
        }
        self.process.stdin.write(json.dumps(init_payload) + "\n")
        self.process.stdin.flush()
        ready = self._read_worker_json("local model startup")
        if not ready.get("ok"):
            raise BranchCleanerError(str(ready.get("error") or "Local model worker failed to start."))

    def _read_worker_json(self, context: str) -> dict[str, object]:
        if self.process is None or self.process.stdout is None:
            raise BranchCleanerError("Local model worker is not running.")
        selector = selectors.DefaultSelector()
        selector.register(self.process.stdout, selectors.EVENT_READ)
        try:
            events = selector.select(self.timeout_seconds)
        finally:
            selector.close()
        if not events:
            self.close(kill=True)
            raise BranchCleanerError(
                f"Timed out waiting for {context} after {self.timeout_seconds}s. "
                "Try a smaller GGUF model, CPU mode, or Gemini mode."
            )
        line = self.process.stdout.readline()
        if not line:
            raise BranchCleanerError(self._worker_failure_message(f"No output from {context}."))
        try:
            payload = json.loads(line)
        except json.JSONDecodeError as exc:
            raise BranchCleanerError(
                f"Local model worker returned non-JSON during {context}: {line[-1000:]}\n"
                f"Worker log:\n{self._stderr_tail()}"
            ) from exc
        if not isinstance(payload, dict):
            raise BranchCleanerError(f"Local model worker returned unexpected payload during {context}.")
        return payload

    def _worker_failure_message(self, prefix: str) -> str:
        return (
            f"{prefix} Exit code: {None if self.process is None else self.process.poll()}.\n"
            "This often means the selected GGUF is incompatible with the installed llama-cpp-python build, "
            "or Metal/GPU execution crashed. Try CPU mode, a smaller Q4 GGUF, or upgrade llama-cpp-python.\n"
            f"Worker log:\n{self._stderr_tail()}"
        )

    def _stderr_tail(self) -> str:
        if self.stderr_handle is not None:
            self.stderr_handle.flush()
        if self.stderr_path is None or not self.stderr_path.exists():
            return "no worker log"
        lines = self.stderr_path.read_text(encoding="utf-8", errors="replace").splitlines()
        return "\n".join(lines[-12:]) or "empty worker log"

    def close(self, *, kill: bool = False) -> None:
        if self.process is not None and self.process.poll() is None:
            if kill:
                self.process.kill()
            else:
                self.process.terminate()
            try:
                self.process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                self.process.kill()
        self.process = None
        if self.stderr_handle is not None:
            self.stderr_handle.close()
            self.stderr_handle = None

    def __del__(self):
        try:
            self.close()
        except Exception:
            pass


LocalModelAddressParser = LlamaCppWorkerAddressParser


def load_llama_cpp_factory():
    try:
        from llama_cpp import Llama
    except ImportError as exc:
        raise BranchCleanerError(
            "Local parsing requires llama-cpp-python. Install it with: "
            "pip install -r requirements.txt"
        ) from exc
    return Llama


def run_llama_completion(model: object, prompt: str) -> str:
    result = model(
        prompt,
        max_tokens=256,
        temperature=0,
        stop=["\n\nAddress:", "\n\nRules:"],
    )
    if isinstance(result, str):
        return result
    if isinstance(result, dict):
        choices = result.get("choices")
        if isinstance(choices, list) and choices:
            first = choices[0]
            if isinstance(first, dict):
                return str(first.get("text") or first.get("message", {}).get("content") or "")
    return str(result or "")


def run_llama_address_chat(model: object, address: str) -> str:
    return run_llama_chat_completion(
        model,
        [
            {
                "role": "system",
                "content": "You are a strict JSON extraction engine. Return only one valid compact JSON object. No markdown.",
            },
            {"role": "user", "content": build_local_llm_prompt(address)},
        ],
    )


def run_llama_repair_chat(model: object, repair_prompt: str) -> str:
    return run_llama_chat_completion(
        model,
        [
            {
                "role": "system",
                "content": "Return only one valid compact JSON object. No markdown, no explanation.",
            },
            {"role": "user", "content": repair_prompt},
        ],
    )


def run_llama_chat_completion(model: object, messages: list[dict[str, str]]) -> str:
    if not hasattr(model, "create_chat_completion"):
        # Test doubles from earlier versions may only implement the raw completion API.
        return run_llama_completion(model, messages[-1]["content"])
    result = model.create_chat_completion(
        messages=messages,
        max_tokens=160,
        temperature=0,
    )
    if isinstance(result, str):
        return result
    if isinstance(result, dict):
        choices = result.get("choices")
        if isinstance(choices, list) and choices:
            first = choices[0]
            if isinstance(first, dict):
                message = first.get("message")
                if isinstance(message, dict) and isinstance(message.get("content"), str):
                    return message["content"]
                if isinstance(first.get("text"), str):
                    return first["text"]
    return str(result or "")


def normalize_spaces(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def is_blank_like(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    text = normalize_spaces(value).lower()
    return text in NULL_LIKE_VALUES


def repo_slug(repo_id: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "__", repo_id).strip("_") or "model"


def format_bytes(size_bytes: int | None) -> str:
    if size_bytes is None:
        return "unknown size"
    units = ["B", "KB", "MB", "GB", "TB"]
    value = float(size_bytes)
    for unit in units:
        if value < 1024 or unit == units[-1]:
            return f"{value:.1f} {unit}" if unit != "B" else f"{int(value)} B"
        value /= 1024
    return f"{size_bytes} B"


def model_index_path(models_dir: Path = MODELS_DIR) -> Path:
    return models_dir / MODEL_INDEX_FILENAME


def read_model_index(models_dir: Path = MODELS_DIR) -> dict[str, object]:
    path = model_index_path(models_dir)
    if not path.exists():
        return {"models": []}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {"models": []}


def write_model_index(index: dict[str, object], models_dir: Path = MODELS_DIR) -> None:
    models_dir.mkdir(parents=True, exist_ok=True)
    model_index_path(models_dir).write_text(json.dumps(index, indent=2, sort_keys=True), encoding="utf-8")


def record_model_in_index(model: LocalGgufModel, models_dir: Path = MODELS_DIR) -> None:
    index = read_model_index(models_dir)
    models = index.setdefault("models", [])
    if not isinstance(models, list):
        models = []
        index["models"] = models
    row = {
        "repo_id": model.repo_id,
        "filename": model.filename,
        "path": str(model.path),
        "size_bytes": model.size_bytes,
    }
    models[:] = [existing for existing in models if not (isinstance(existing, dict) and existing.get("path") == str(model.path))]
    models.append(row)
    write_model_index(index, models_dir)


def discover_local_gguf_models(models_dir: Path = MODELS_DIR) -> list[LocalGgufModel]:
    index = read_model_index(models_dir)
    metadata_by_path: dict[str, dict[str, object]] = {}
    for row in index.get("models", []) if isinstance(index.get("models"), list) else []:
        if isinstance(row, dict) and row.get("path"):
            metadata_by_path[str(row["path"])] = row
    discovered: list[LocalGgufModel] = []
    if not models_dir.exists():
        return discovered
    for path in sorted(models_dir.rglob("*.gguf")):
        metadata = metadata_by_path.get(str(path), {})
        filename = str(metadata.get("filename") or path.name)
        repo_id = str(metadata.get("repo_id") or path.parent.name.replace("__", "/"))
        size_value = metadata.get("size_bytes")
        size_bytes = int(size_value) if isinstance(size_value, (int, float)) else path.stat().st_size
        discovered.append(LocalGgufModel(repo_id=repo_id, filename=filename, path=path, size_bytes=size_bytes))
    return discovered


def search_huggingface_gguf_models(query: str, *, limit: int = 8, api: object | None = None) -> list[HuggingFaceGgufFile]:
    if api is None:
        try:
            from huggingface_hub import HfApi
        except ImportError as exc:
            raise BranchCleanerError(
                "Hugging Face search requires huggingface_hub. Install it with: "
                "pip install -r requirements.txt"
            ) from exc
        api = HfApi()
    try:
        models = list(api.list_models(search=query, filter="gguf", sort="downloads", direction=-1, limit=limit))
    except TypeError:
        models = list(api.list_models(search=query, limit=limit))
    results: list[HuggingFaceGgufFile] = []
    for model_info in models:
        repo_id = getattr(model_info, "modelId", None) or getattr(model_info, "id", None)
        if not repo_id:
            continue
        siblings = getattr(model_info, "siblings", None)
        if not siblings and hasattr(api, "model_info"):
            try:
                siblings = getattr(api.model_info(repo_id), "siblings", None)
            except Exception:
                siblings = []
        for sibling in siblings or []:
            filename = getattr(sibling, "rfilename", None) or getattr(sibling, "filename", None)
            if not filename or not str(filename).lower().endswith(".gguf"):
                continue
            size = getattr(sibling, "size", None)
            results.append(HuggingFaceGgufFile(repo_id=str(repo_id), filename=str(filename), size_bytes=size))
    return prioritize_gguf_results(results)[:limit]


def prioritize_gguf_results(results: list[HuggingFaceGgufFile]) -> list[HuggingFaceGgufFile]:
    preferred = ("q4_k_m", "q4_0", "q3_k_m", "q2_k")

    def score(result: HuggingFaceGgufFile) -> tuple[int, int, str]:
        filename = result.filename.lower()
        quant_score = next((index for index, token in enumerate(preferred) if token in filename), len(preferred))
        size_score = result.size_bytes if isinstance(result.size_bytes, int) else LOCAL_MODEL_TOO_LARGE_BYTES
        return quant_score, size_score, f"{result.repo_id}/{result.filename}"

    return sorted(results, key=score)


def download_hf_gguf_model(
    selection: HuggingFaceGgufFile,
    *,
    models_dir: Path = MODELS_DIR,
    token: str | None = None,
    downloader: Callable[..., str] | None = None,
) -> LocalGgufModel:
    if downloader is None:
        try:
            from huggingface_hub import hf_hub_download
        except ImportError as exc:
            raise BranchCleanerError(
                "Model download requires huggingface_hub. Install it with: "
                "pip install -r requirements.txt"
            ) from exc
        downloader = hf_hub_download
    local_dir = models_dir / repo_slug(selection.repo_id)
    local_dir.mkdir(parents=True, exist_ok=True)
    try:
        downloaded_path = Path(
            downloader(
                repo_id=selection.repo_id,
                filename=selection.filename,
                local_dir=str(local_dir),
                token=token,
            )
        )
    except Exception as exc:
        if token:
            raise
        hf_token = getpass.getpass("Download failed. Hugging Face token for gated/private model (blank to cancel): ").strip()
        if not hf_token:
            raise BranchCleanerError(f"Could not download model: {exc}") from exc
        downloaded_path = Path(
            downloader(
                repo_id=selection.repo_id,
                filename=selection.filename,
                local_dir=str(local_dir),
                token=hf_token,
            )
        )
    local_model = LocalGgufModel(
        repo_id=selection.repo_id,
        filename=selection.filename,
        path=downloaded_path,
        size_bytes=selection.size_bytes if selection.size_bytes is not None else downloaded_path.stat().st_size,
    )
    record_model_in_index(local_model, models_dir)
    return local_model


def parse_coordinate(value: object, coordinate_kind: str) -> float | None:
    if is_blank_like(value):
        return None
    if isinstance(value, bool):
        return None
    try:
        number = float(str(value).strip())
    except (TypeError, ValueError):
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    if coordinate_kind == LATITUDE_COLUMN and not -90 <= number <= 90:
        return None
    if coordinate_kind == LONGITUDE_COLUMN and not -180 <= number <= 180:
        return None
    return number


def has_valid_coordinates(row_values: dict[str, object]) -> bool:
    return (
        parse_coordinate(row_values.get(LATITUDE_COLUMN), LATITUDE_COLUMN) is not None
        and parse_coordinate(row_values.get(LONGITUDE_COLUMN), LONGITUDE_COLUMN) is not None
    )


def has_valid_selected_coordinates(sheet: Worksheet, row_number: int, columns: CoordinateColumns) -> bool:
    return (
        parse_coordinate(sheet.cell(row=row_number, column=columns.latitude).value, LATITUDE_COLUMN) is not None
        and parse_coordinate(sheet.cell(row=row_number, column=columns.longitude).value, LONGITUDE_COLUMN) is not None
    )


def find_candidate_workbooks(directory: Path) -> list[Path]:
    candidates = []
    for path in sorted([*directory.glob("*.xlsx"), *directory.glob("*.csv")]):
        name_lower = path.name.lower()
        if path.name.startswith("~$"):
            continue
        if re.search(r"_cleaned(?:_\d+)?\.(?:xlsx|csv)$", name_lower):
            continue
        candidates.append(path)
    return candidates


def unique_cleaned_output_path(source_path: Path) -> Path:
    base = source_path.with_name(f"{source_path.stem}_cleaned{source_path.suffix}")
    if not base.exists():
        return base
    counter = 2
    while True:
        candidate = source_path.with_name(f"{source_path.stem}_cleaned_{counter}{source_path.suffix}")
        if not candidate.exists():
            return candidate
        counter += 1


def load_workbook_for_cleaning(workbook_path: Path):
    return load_workbook(workbook_path)


def load_csv_for_cleaning(csv_path: Path) -> LoadedFile:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = csv_path.stem[:31] or "csv"
    with csv_path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        for row in reader:
            sheet.append(row)
    if sheet.max_row < 1:
        raise BranchCleanerError(f"CSV file is empty: {csv_path}")
    return LoadedFile(path=csv_path, kind="csv", workbook=workbook, sheet=sheet)


def load_file_for_cleaning(path: Path) -> LoadedFile:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return load_csv_for_cleaning(path)
    if suffix == ".xlsx":
        workbook = load_workbook(path)
        sheet = select_default_sheet(workbook)
        return LoadedFile(path=path, kind="xlsx", workbook=workbook, sheet=sheet)
    raise BranchCleanerError(f"Unsupported file type: {path.suffix}")


def save_loaded_file(loaded_file: LoadedFile, output_path: Path) -> None:
    if loaded_file.kind == "csv":
        save_sheet_as_csv(loaded_file.sheet, output_path)
        return
    loaded_file.workbook.save(output_path)


def save_sheet_as_csv(sheet: Worksheet, output_path: Path) -> None:
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if value is None else value for value in row])


def load_branch_sheet(workbook_path: Path):
    workbook = load_workbook(workbook_path)
    if BRANCH_SHEET not in workbook.sheetnames:
        raise BranchCleanerError(f"Workbook does not contain a '{BRANCH_SHEET}' sheet: {workbook_path}")
    return workbook, workbook[BRANCH_SHEET]


def select_default_sheet(workbook) -> Worksheet:
    if BRANCH_SHEET in workbook.sheetnames:
        print(f"Selected sheet: {BRANCH_SHEET}")
        return workbook[BRANCH_SHEET]
    return prompt_sheet_selection(workbook)


def headers_in_order(sheet: Worksheet) -> list[HeaderColumn]:
    headers: list[HeaderColumn] = []
    for cell in sheet[1]:
        label = normalize_spaces(cell.value)
        if label:
            headers.append(HeaderColumn(label=label, index=cell.column))
    return headers


def header_map(sheet: Worksheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in sheet[1]:
        if cell.value is None:
            continue
        name = normalize_spaces(cell.value)
        if name:
            headers[name] = cell.column
    return headers


def last_header_column(sheet: Worksheet) -> int:
    last_column = 0
    for cell in sheet[1]:
        if normalize_spaces(cell.value):
            last_column = max(last_column, cell.column)
    return last_column


def append_fresh_columns(sheet: Worksheet, columns: list[str]) -> dict[str, int]:
    next_column = last_header_column(sheet) + 1
    appended: dict[str, int] = {}
    for column in columns:
        sheet.cell(row=1, column=next_column).value = column
        appended[column] = next_column
        next_column += 1
    return appended


def validate_column_index(sheet: Worksheet, column_index: int, purpose: str) -> int:
    if column_index < 1 or column_index > max(sheet.max_column, 1):
        raise BranchCleanerError(f"Selected {purpose} column is outside the worksheet range: {column_index}")
    if not normalize_spaces(sheet.cell(row=1, column=column_index).value):
        raise BranchCleanerError(f"Selected {purpose} column has no header: column {column_index}")
    return column_index


def require_columns(sheet: Worksheet, columns: list[str]) -> dict[str, int]:
    headers = header_map(sheet)
    missing = [column for column in columns if column not in headers]
    if missing:
        raise BranchCleanerError(f"Missing required column(s) in '{BRANCH_SHEET}': {', '.join(missing)}")
    return headers


def ensure_columns(sheet: Worksheet, columns: list[str]) -> dict[str, int]:
    headers = header_map(sheet)
    next_column = sheet.max_column + 1
    for column in columns:
        if column not in headers:
            sheet.cell(row=1, column=next_column).value = column
            headers[column] = next_column
            next_column += 1
    return headers


def row_dict(sheet: Worksheet, row_number: int, headers: dict[str, int]) -> dict[str, object]:
    return {name: sheet.cell(row=row_number, column=column).value for name, column in headers.items()}


def infer_default_coordinate_columns(sheet: Worksheet) -> CoordinateColumns:
    headers = header_map(sheet)
    missing = [column for column in [LATITUDE_COLUMN, LONGITUDE_COLUMN, FULL_ADDRESS_COLUMN] if column not in headers]
    if missing:
        raise BranchCleanerError(f"Missing required column(s) in '{sheet.title}': {', '.join(missing)}")
    return CoordinateColumns(
        latitude=headers[LATITUDE_COLUMN],
        longitude=headers[LONGITUDE_COLUMN],
        address=headers[FULL_ADDRESS_COLUMN],
    )


def inspect_coordinate_needs(
    sheet: Worksheet,
    columns: CoordinateColumns | None = None,
) -> tuple[CoordinateColumns, CoordinateCleanSummary, list[tuple[int, str]]]:
    columns = columns or infer_default_coordinate_columns(sheet)
    validate_column_index(sheet, columns.latitude, "latitude")
    validate_column_index(sheet, columns.longitude, "longitude")
    validate_column_index(sheet, columns.address, "address")
    rows_to_geocode: list[tuple[int, str]] = []
    summary = CoordinateCleanSummary(rows_checked=max(sheet.max_row - 1, 0))

    for row_number in range(2, sheet.max_row + 1):
        if has_valid_selected_coordinates(sheet, row_number, columns):
            continue
        summary.rows_needing_coordinates += 1
        address = normalize_spaces(sheet.cell(row=row_number, column=columns.address).value)
        if not address:
            summary.skipped_no_address_rows.append(row_number)
            continue
        rows_to_geocode.append((row_number, address))

    return columns, summary, rows_to_geocode


def clean_missing_coordinates(
    sheet: Worksheet,
    geocoder: Geocoder | None = None,
    columns: CoordinateColumns | None = None,
) -> CoordinateCleanSummary:
    columns, summary, rows_to_geocode = inspect_coordinate_needs(sheet, columns)
    if not rows_to_geocode:
        return summary
    if geocoder is None:
        raise BranchCleanerError("Rows need coordinates, but no geocoder was provided.")

    for row_number, address in rows_to_geocode:
        try:
            result = geocoder.geocode(address)
        except Exception as exc:  # noqa: BLE001 - keep cleaning remaining rows after one API failure.
            summary.failed_rows[row_number] = str(exc)
            continue
        if result is None:
            summary.no_result_rows.append(row_number)
            continue
        latitude, longitude = result
        sheet.cell(row=row_number, column=columns.latitude).value = latitude
        sheet.cell(row=row_number, column=columns.longitude).value = longitude
        summary.geocoded_rows += 1

    return summary


def parse_full_addresses(
    sheet: Worksheet,
    parser: AddressParser,
    address_column: int | None = None,
) -> ParseAddressSummary:
    if address_column is None:
        headers = require_columns(sheet, [FULL_ADDRESS_COLUMN])
        address_column = headers[FULL_ADDRESS_COLUMN]
    validate_column_index(sheet, address_column, "address")
    output_columns = append_fresh_columns(sheet, PARSED_ADDRESS_COLUMNS)
    summary = ParseAddressSummary(rows_checked=max(sheet.max_row - 1, 0))

    for row_number in range(2, sheet.max_row + 1):
        address = normalize_spaces(sheet.cell(row=row_number, column=address_column).value)
        if not address:
            summary.skipped_no_address_rows.append(row_number)
            write_parsed_address(sheet, row_number, output_columns, empty_address_result())
            continue
        try:
            parsed = parser.parse_address(address)
            normalized = normalize_address_parse_result(parsed)
        except Exception as exc:  # noqa: BLE001 - keep cleaning remaining rows after one parser failure.
            summary.failed_rows[row_number] = str(exc)
            normalized = empty_address_result()
        write_parsed_address(sheet, row_number, output_columns, normalized)
        if row_number not in summary.failed_rows:
            summary.parsed_rows += 1

    return summary


def write_parsed_address(sheet: Worksheet, row_number: int, output_columns: dict[str, int], parsed: dict[str, str | None]) -> None:
    for source_field, output_column in zip(ADDRESS_FIELDS, PARSED_ADDRESS_COLUMNS, strict=True):
        sheet.cell(row=row_number, column=output_columns[output_column]).value = parsed.get(source_field)


def empty_address_result() -> dict[str, None]:
    return {field: None for field in ADDRESS_FIELDS}


def normalize_address_parse_result(raw: object) -> dict[str, str | None]:
    if isinstance(raw, str):
        raw = json.loads(extract_json_object_text(raw))
    if not isinstance(raw, dict):
        raise BranchCleanerError("Address parser returned a non-object JSON value.")

    aliases = {
        "barangay": "barangay",
        "city_municipality": "city_municipality",
        "city/municipality": "city_municipality",
        "city": "city_municipality",
        "municipality": "city_municipality",
        "province": "province",
        "region": "region",
        "zip_code": "zip_code",
        "zip code": "zip_code",
        "zipcode": "zip_code",
        "postal_code": "zip_code",
        "country": "country",
        "island_group": "island_group",
        "island group": "island_group",
        "islandgroup": "island_group",
    }
    normalized = empty_address_result()
    for key, value in raw.items():
        canonical = aliases.get(normalize_spaces(key).lower())
        if canonical not in normalized:
            continue
        text = normalize_spaces(value)
        if canonical == "island_group":
            normalized[canonical] = normalize_island_group(text)
            continue
        normalized[canonical] = None if is_blank_like(text) else text
    return normalized


def normalize_island_group(value: object) -> str | None:
    text = normalize_spaces(value).lower()
    if is_blank_like(text):
        return None
    if "luzon" in text:
        return "Luzon"
    if "visayas" in text or "visaya" in text:
        return "Visayas"
    if "mindanao" in text:
        return "Mindanao"
    return None


def extract_json_object_text(text: str) -> str:
    stripped = text.strip()
    if stripped.startswith("```"):
        stripped = re.sub(r"^```(?:json)?\s*", "", stripped, flags=re.I)
        stripped = re.sub(r"\s*```$", "", stripped)
    start = stripped.find("{")
    end = stripped.rfind("}")
    if start < 0 or end < start:
        raise BranchCleanerError("No JSON object found in parser response.")
    return stripped[start : end + 1]


def extract_gemini_json_payload(payload: object) -> object:
    if isinstance(payload, dict) and all(field in payload for field in ADDRESS_FIELDS):
        return payload
    if not isinstance(payload, dict):
        raise BranchCleanerError("Gemini API returned an unexpected response.")

    for key in ("output_text", "outputText", "text", "content"):
        value = payload.get(key)
        if isinstance(value, str) and value.strip():
            return json.loads(extract_json_object_text(value))

    candidates = payload.get("candidates")
    if isinstance(candidates, list):
        for candidate in candidates:
            if not isinstance(candidate, dict):
                continue
            parts = candidate.get("content", {}).get("parts", [])
            if not isinstance(parts, list):
                continue
            for part in parts:
                if isinstance(part, dict) and isinstance(part.get("text"), str):
                    return json.loads(extract_json_object_text(part["text"]))

    # Defensive support for response shapes that expose text in step/content arrays.
    possible_texts: list[str] = []
    for step in payload.get("steps", []) or []:
        if isinstance(step, dict):
            for key in ("text", "output_text", "outputText"):
                value = step.get(key)
                if isinstance(value, str):
                    possible_texts.append(value)
    for candidate in possible_texts:
        try:
            return json.loads(extract_json_object_text(candidate))
        except Exception:
            continue

    raise BranchCleanerError("Could not extract JSON text from Gemini API response.")


def address_json_schema() -> dict[str, object]:
    nullable_string = {"anyOf": [{"type": "string"}, {"type": "null"}]}
    return {
        "type": "object",
        "properties": {field: nullable_string for field in ADDRESS_FIELDS},
        "required": ADDRESS_FIELDS,
    }


def build_address_prompt(address: str) -> str:
    return f"""Parse the given Philippine address into structured location fields.

Extract only these fields:
- Barangay
- City/Municipality
- Province
- Region (For region, return the official Philippine region code format only, such as "Region I", "Region II", "Region III", "Region IV-A", "Region VI", etc. Do not return the region name unless separately requested.)
- ZIP Code
- Country
- Island Group

Rules:
- Return only valid JSON.
- If a field is missing or cannot be confidently determined, use null.
- Normalize abbreviations:
  - "Brgy." = Barangay
  - "Sta." = Santa
- Do not include building name, street, sitio, purok, subdivision, or landmark.
- Country should be "Philippines" if the address is in the Philippines.
- Island Group must be exactly one of "Luzon", "Visayas", or "Mindanao"; use null if it cannot be confidently determined from the address.
- Do not guess ZIP Code unless it is clearly present in the address.
- If the address contains an inconsistency, prioritize barangay-to-city/municipality accuracy over the written city/province.

Address:
{address}

JSON format:
{{
  "barangay": null,
  "city_municipality": null,
  "province": null,
  "region": null,
  "zip_code": null,
  "country": null,
  "island_group": null
}}"""


def build_local_llm_prompt(address: str) -> str:
    return f"""You parse Philippine addresses into administrative location fields.
Return only JSON with keys: barangay, city_municipality, province, region, zip_code, country, island_group.

Rules:
- Building names, company names, malls, roads, streets, drives, highways, subdivisions, landmarks, and branch names are not barangays, cities, or provinces.
- If barangay is not explicitly written as Brgy/Barangay or clearly known, use null.
- Never copy city_municipality into barangay unless the address explicitly marks that same value as Brgy/Barangay.
- ZIP code must be null unless a numeric postal code appears in the address.
- For Metro Manila cities, province is "Metro Manila" and region is "NCR".
- Island group must be exactly "Luzon", "Visayas", or "Mindanao"; use null if not confidently known.
- Do not put road names, building names, or ZIP/postal concepts into city_municipality.
- Country should be "Philippines" when the address is in the Philippines.

Examples:
Address: Ropali Plaza, Escriva Drive, Pasig, Metro Manila
{{"barangay":null,"city_municipality":"Pasig","province":"Metro Manila","region":"NCR","zip_code":null,"country":"Philippines","island_group":"Luzon"}}

Address: 123 Brgy. San Antonio, Makati City, Metro Manila 1203
{{"barangay":"Barangay San Antonio","city_municipality":"Makati City","province":"Metro Manila","region":"NCR","zip_code":"1203","country":"Philippines","island_group":"Luzon"}}

Address: National Highway, Barangay San Miguel, Calasiao, Pangasinan
{{"barangay":"Barangay San Miguel","city_municipality":"Calasiao","province":"Pangasinan","region":"Region I","zip_code":null,"country":"Philippines","island_group":"Luzon"}}

Now parse this address:
{address}"""


def build_local_json_repair_prompt(address: str, previous_output: str) -> str:
    return f"""The previous response was not valid JSON for the required address parsing task.

Return only one valid JSON object with exactly these keys:
barangay, city_municipality, province, region, zip_code, country, island_group.
Use null for missing values.
island_group must be exactly Luzon, Visayas, or Mindanao when known.

Address:
{address}

Invalid previous response:
{previous_output}
"""


def prompt_workbook_selection(workbooks: list[Path]) -> Path:
    if not workbooks:
        raise BranchCleanerError("No eligible .xlsx or .csv files found beside this script.")
    if len(workbooks) == 1:
        print(f"Selected workbook: {workbooks[0].name}")
        return workbooks[0]

    print("Available workbooks:")
    for index, path in enumerate(workbooks, start=1):
        print(f"  {index}. {path.name}")
    while True:
        choice = input("Choose workbook number: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(workbooks):
            return workbooks[int(choice) - 1]
        print("Please enter a valid workbook number.")


def prompt_sheet_selection(workbook) -> Worksheet:
    print("Available sheets:")
    for index, sheet_name in enumerate(workbook.sheetnames, start=1):
        print(f"  {index}. {sheet_name}")
    while True:
        choice = input("Choose sheet number: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(workbook.sheetnames):
            sheet = workbook[workbook.sheetnames[int(choice) - 1]]
            print(f"Selected sheet: {sheet.title}")
            return sheet
        print("Please enter a valid sheet number.")


def prompt_column_selection(sheet: Worksheet, purpose: str) -> int:
    headers = headers_in_order(sheet)
    if not headers:
        raise BranchCleanerError(f"Worksheet '{sheet.title}' does not have any headers in row 1.")
    print(f"Available columns for {purpose}:")
    for display_index, header in enumerate(headers, start=1):
        print(f"  {display_index}. {header.label} (column {header.index})")
    while True:
        choice = input(f"Choose {purpose} column number: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(headers):
            selected = headers[int(choice) - 1]
            print(f"Selected {purpose} column: {selected.label}")
            return selected.index
        print("Please enter a valid column number.")


def prompt_actions() -> set[str]:
    print("Cleaning options:")
    print("  1. No coordinates")
    print("  2. Parse address with LLM")
    print("  3. Both coordinates and LLM address parsing")
    while True:
        choice = input("Choose cleaning option [1/2/3]: ").strip()
        if choice == "1":
            return {"coordinates"}
        if choice == "2":
            return {"parse_address"}
        if choice == "3":
            return {"coordinates", "parse_address"}
        print("Please choose 1, 2, or 3.")


def build_parser_from_prompt() -> AddressParser:
    provider = input("Address parser provider [gemini/local] (default: gemini): ").strip().lower() or "gemini"
    if provider == "local":
        return build_local_parser_from_prompt()
    if provider != "gemini":
        raise BranchCleanerError(f"Unsupported parser provider: {provider}")
    api_key = getpass.getpass("Gemini API key (not stored): ").strip()
    if not api_key:
        raise BranchCleanerError("Gemini API key is required for address parsing.")
    model = input(f"Gemini model [{DEFAULT_GEMINI_MODEL}]: ").strip() or DEFAULT_GEMINI_MODEL
    return GeminiApiAddressParser(api_key=api_key, model=model)


def build_local_parser_from_prompt(models_dir: Path = MODELS_DIR) -> AddressParser:
    selected_model = prompt_local_model(models_dir)
    if selected_model.size_bytes and selected_model.size_bytes > LOCAL_MODEL_TOO_LARGE_BYTES:
        print(
            f"Warning: selected model is {format_bytes(selected_model.size_bytes)}. "
            "This may be slow or fail on an 8 GB machine."
        )
    n_gpu_layers, flash_attn = prompt_local_runtime()
    print(f"Selected local model: {selected_model.path}")
    print("Local parsing will run in an isolated worker process so native llama.cpp crashes do not stop the cleaner.")
    return LlamaCppWorkerAddressParser(
        selected_model.path,
        n_gpu_layers=n_gpu_layers,
        flash_attn=flash_attn,
        n_threads=LOCAL_DEFAULT_THREADS,
        n_threads_batch=LOCAL_DEFAULT_THREADS,
    )


def prompt_local_runtime() -> tuple[int, bool]:
    print("llama.cpp runtime:")
    print("  GPU layers 0 = CPU only, usually safest.")
    print("  GPU layers -1 = offload all possible layers to Metal/GPU if your GGUF/runtime supports it.")
    print("  Flash Attention can help some models but may be unsupported or unstable for others.")
    while True:
        value = input("GPU layers [0]: ").strip() or "0"
        try:
            gpu_layers = int(value)
            break
        except ValueError:
            print("Please enter an integer, for example 0 or -1.")
    flash_attn = prompt_yes_no("Enable Flash Attention", False)
    return gpu_layers, flash_attn


def prompt_yes_no(label: str, default: bool) -> bool:
    suffix = "Y/n" if default else "y/N"
    while True:
        value = input(f"{label} [{suffix}]: ").strip().lower()
        if not value:
            return default
        if value in {"y", "yes"}:
            return True
        if value in {"n", "no"}:
            return False
        print("Please enter y or n.")


def prompt_local_model(models_dir: Path = MODELS_DIR) -> LocalGgufModel:
    while True:
        local_models = discover_local_gguf_models(models_dir)
        print("Local model options:")
        print("  1. Use downloaded model")
        print("  2. Search Hugging Face")
        print("  3. Enter Hugging Face repo/file manually")
        choice = input("Choose local model option [1/2/3]: ").strip()
        if choice == "1":
            if not local_models:
                print("No downloaded .gguf models found yet.")
                continue
            return prompt_existing_local_model(local_models)
        if choice == "2":
            return search_and_download_model(models_dir)
        if choice == "3":
            return prompt_manual_hf_model(models_dir)
        print("Please choose 1, 2, or 3.")


def prompt_existing_local_model(local_models: list[LocalGgufModel]) -> LocalGgufModel:
    print("Downloaded GGUF models:")
    for index, model in enumerate(local_models, start=1):
        print(f"  {index}. {model.repo_id} / {model.filename} ({format_bytes(model.size_bytes)})")
    while True:
        choice = input("Choose downloaded model number: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(local_models):
            return local_models[int(choice) - 1]
        print("Please enter a valid model number.")


def search_and_download_model(models_dir: Path = MODELS_DIR) -> LocalGgufModel:
    query = input("Hugging Face search query [GGUF instruct small]: ").strip() or "GGUF instruct small"
    results = search_huggingface_gguf_models(query)
    if not results:
        raise BranchCleanerError(f"No GGUF model files found for search query: {query}")
    print("Hugging Face GGUF results:")
    for index, result in enumerate(results, start=1):
        warning = " - large for 8 GB RAM" if result.size_bytes and result.size_bytes > LOCAL_MODEL_TOO_LARGE_BYTES else ""
        print(f"  {index}. {result.repo_id} / {result.filename} ({format_bytes(result.size_bytes)}){warning}")
    while True:
        choice = input("Choose model file number to download: ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(results):
            return download_hf_gguf_model(results[int(choice) - 1], models_dir=models_dir)
        print("Please enter a valid model number.")


def prompt_manual_hf_model(models_dir: Path = MODELS_DIR) -> LocalGgufModel:
    repo_id = input("Hugging Face repo id, e.g. owner/model-GGUF: ").strip()
    filename = input("GGUF filename in that repo: ").strip()
    if not repo_id or not filename:
        raise BranchCleanerError("Both Hugging Face repo id and GGUF filename are required.")
    if not filename.lower().endswith(".gguf"):
        raise BranchCleanerError("Local model filename must end with .gguf")
    return download_hf_gguf_model(HuggingFaceGgufFile(repo_id=repo_id, filename=filename), models_dir=models_dir)


def print_coordinate_summary(summary: CoordinateCleanSummary) -> None:
    print("\nCoordinate cleaning summary")
    print(f"  Rows checked: {summary.rows_checked}")
    print(f"  Rows needing coordinates: {summary.rows_needing_coordinates}")
    print(f"  Rows geocoded: {summary.geocoded_rows}")
    print(f"  Skipped no address: {len(summary.skipped_no_address_rows)}")
    print(f"  No geocode result: {len(summary.no_result_rows)}")
    print(f"  Failed rows: {len(summary.failed_rows)}")
    if summary.rows_needing_coordinates == 0:
        print("  All rows already had valid coordinates. Existing coordinates were not overwritten.")


def print_parse_summary(summary: ParseAddressSummary) -> None:
    print("\nAddress parse summary")
    print(f"  Rows checked: {summary.rows_checked}")
    print(f"  Rows parsed: {summary.parsed_rows}")
    print(f"  Skipped no full_address: {len(summary.skipped_no_address_rows)}")
    print(f"  Failed rows: {len(summary.failed_rows)}")
    for row_number, error in list(summary.failed_rows.items())[:3]:
        print(f"  Row {row_number} error: {error}")
    if len(summary.failed_rows) > 3:
        print(f"  ...and {len(summary.failed_rows) - 3} more failed rows.")


def run_cleaner(script_directory: Path) -> Path:
    workbooks = find_candidate_workbooks(script_directory)
    workbook_path = prompt_workbook_selection(workbooks)
    actions = prompt_actions()
    loaded_file = load_file_for_cleaning(workbook_path)
    sheet = loaded_file.sheet

    if "coordinates" in actions:
        coordinate_columns = CoordinateColumns(
            latitude=prompt_column_selection(sheet, "latitude"),
            longitude=prompt_column_selection(sheet, "longitude"),
            address=prompt_column_selection(sheet, "address for geocoding"),
        )
        _, preliminary, rows_to_geocode = inspect_coordinate_needs(sheet, coordinate_columns)
        if not rows_to_geocode:
            print_coordinate_summary(preliminary)
        else:
            api_key = getpass.getpass("Google Geocoding API key (not stored): ").strip()
            if not api_key:
                raise BranchCleanerError("Google Geocoding API key is required for coordinate cleaning.")
            summary = clean_missing_coordinates(sheet, geocoder=GoogleGeocoder(api_key), columns=coordinate_columns)
            print_coordinate_summary(summary)

    if "parse_address" in actions:
        address_column = prompt_column_selection(sheet, "address for parsing")
        parser = build_parser_from_prompt()
        summary = parse_full_addresses(sheet, parser, address_column=address_column)
        print_parse_summary(summary)

    output_path = unique_cleaned_output_path(workbook_path)
    save_loaded_file(loaded_file, output_path)
    print(f"\nSaved cleaned file: {output_path.name}")
    return output_path


def local_parse_worker_main() -> int:
    try:
        payload = json.loads(sys.stdin.read() or "{}")
        model_path = Path(str(payload["model_path"]))
        address = str(payload["address"])
        parser = LlamaCppAddressParser(
            model_path,
            n_ctx=int(payload.get("n_ctx") or 1536),
            n_gpu_layers=int(payload.get("n_gpu_layers") if payload.get("n_gpu_layers") is not None else -1),
            n_batch=int(payload.get("n_batch") or 512),
            n_ubatch=int(payload.get("n_ubatch") or 128),
            n_threads=payload.get("n_threads"),
            n_threads_batch=payload.get("n_threads_batch"),
            flash_attn=bool(payload.get("flash_attn")),
        )
        parsed = parser.parse_address(address)
        print(json.dumps({"ok": True, "parsed": parsed}, ensure_ascii=False))
        return 0
    except Exception as exc:  # noqa: BLE001 - worker must serialize any recoverable failure for the parent.
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        return 1


def local_parse_server_main() -> int:
    try:
        init_line = sys.stdin.readline()
        init_payload = json.loads(init_line or "{}")
        model_path = Path(str(init_payload["model_path"]))
        parser = LlamaCppAddressParser(
            model_path,
            n_ctx=int(init_payload.get("n_ctx") or 1536),
            n_gpu_layers=int(init_payload.get("n_gpu_layers") if init_payload.get("n_gpu_layers") is not None else -1),
            n_batch=int(init_payload.get("n_batch") or 512),
            n_ubatch=int(init_payload.get("n_ubatch") or 128),
            n_threads=init_payload.get("n_threads"),
            n_threads_batch=init_payload.get("n_threads_batch"),
            flash_attn=bool(init_payload.get("flash_attn")),
        )
        print(json.dumps({"ok": True, "ready": True}), flush=True)

        for line in sys.stdin:
            try:
                request = json.loads(line or "{}")
                parsed = parser.parse_address(str(request.get("address") or ""))
                print(json.dumps({"ok": True, "parsed": parsed}, ensure_ascii=False), flush=True)
            except Exception as exc:  # noqa: BLE001 - keep the local worker alive after row-level failures.
                print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), flush=True)
        return 0
    except Exception as exc:  # noqa: BLE001 - startup failure must be serialized for the parent when possible.
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), flush=True)
        return 1


def main() -> int:
    if len(sys.argv) > 1 and sys.argv[1] == "--local-parse-worker":
        return local_parse_worker_main()
    if len(sys.argv) > 1 and sys.argv[1] == "--local-parse-server":
        return local_parse_server_main()
    work_directory = Path(os.getenv("BRANCH_CLEANER_WORKDIR", Path.cwd())).resolve()
    if len(sys.argv) > 1:
        if len(sys.argv) == 3 and sys.argv[1] == "--directory":
            work_directory = Path(sys.argv[2]).expanduser().resolve()
        else:
            print("Usage: python -m branch_cleaner.branch_sheet_cleaner [--directory PATH]", file=sys.stderr)
            return 2
    try:
        run_cleaner(work_directory)
        return 0
    except BranchCleanerError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
