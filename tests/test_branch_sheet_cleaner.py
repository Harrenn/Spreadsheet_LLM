from __future__ import annotations

import contextlib
import csv
import io
import json
import tempfile
import unittest
from dataclasses import dataclass
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from branch_cleaner import branch_sheet_cleaner as cleaner


class FakeGeocoder:
    def __init__(self):
        self.addresses: list[str] = []

    def geocode(self, address: str):
        self.addresses.append(address)
        if "No Result" in address:
            return None
        return 14.5995, 120.9842


class FakeAddressParser:
    def __init__(self):
        self.addresses: list[str] = []

    def parse_address(self, address: str):
        self.addresses.append(address)
        return {
            "barangay": "Barangay San Antonio",
            "city_municipality": "Pasig",
            "province": "Metro Manila",
            "region": "NCR",
            "zip_code": None,
            "country": "Philippines",
            "island_group": "Luzon",
        }


class FakeLlama:
    def __init__(self, outputs: list[str]):
        self.outputs = outputs
        self.prompts: list[str] = []

    def __call__(self, prompt: str, **kwargs):
        self.prompts.append(prompt)
        text = self.outputs.pop(0)
        return {"choices": [{"text": text}]}


class FakeChatLlama:
    def __init__(self, outputs: list[str]):
        self.outputs = outputs
        self.messages: list[list[dict[str, str]]] = []

    def create_chat_completion(self, messages, **kwargs):
        self.messages.append(messages)
        return {"choices": [{"message": {"content": self.outputs.pop(0)}}]}


@dataclass
class FakeSibling:
    rfilename: str
    size: int | None = None


@dataclass
class FakeModelInfo:
    modelId: str
    siblings: list[FakeSibling]


class FakeHfApi:
    def list_models(self, **kwargs):
        return [
            FakeModelInfo(
                modelId="test/repo-GGUF",
                siblings=[
                    FakeSibling("model-q8_0.gguf", 5_000_000_000),
                    FakeSibling("model-q4_k_m.gguf", 700_000_000),
                    FakeSibling("README.md", 100),
                ],
            )
        ]


def json_payload(city: str) -> str:
    return json.dumps(
        {
            "ok": True,
            "parsed": {
                "barangay": None,
                "city_municipality": city,
                "province": "Metro Manila",
                "region": "NCR",
                "zip_code": None,
                "country": "Philippines",
                "island_group": "Luzon",
            },
        }
    )


def make_workbook(headers: list[str], rows: list[list[object]]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "branch"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    other = workbook.create_sheet("other_sheet")
    other.append(["kept"])
    other.append(["yes"])
    return workbook


class BranchSheetCleanerTests(unittest.TestCase):
    def test_find_candidate_workbooks_excludes_temp_and_cleaned_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            for name in ["dealer.xlsx", "branches.csv", "~$dealer.xlsx", "dealer_cleaned.xlsx", "branches_cleaned.csv", "dealer_cleaned_2.xlsx"]:
                (directory / name).touch()

            candidates = cleaner.find_candidate_workbooks(directory)

            self.assertEqual(candidates, [directory / "branches.csv", directory / "dealer.xlsx"])

    def test_unique_cleaned_output_path_avoids_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "dealer.xlsx"
            source.touch()
            (Path(tmp) / "dealer_cleaned.xlsx").touch()

            output = cleaner.unique_cleaned_output_path(source)

            self.assertEqual(output.name, "dealer_cleaned_2.xlsx")

    def test_coordinate_validation_detects_blank_null_text_and_partial_values(self) -> None:
        self.assertFalse(cleaner.has_valid_coordinates({"latitude": "", "longitude": "120.1"}))
        self.assertFalse(cleaner.has_valid_coordinates({"latitude": "null", "longitude": "120.1"}))
        self.assertFalse(cleaner.has_valid_coordinates({"latitude": "nan", "longitude": "120.1"}))
        self.assertFalse(cleaner.has_valid_coordinates({"latitude": "Pasig", "longitude": "120.1"}))
        self.assertFalse(cleaner.has_valid_coordinates({"latitude": "14.5", "longitude": ""}))
        self.assertTrue(cleaner.has_valid_coordinates({"latitude": "14.5", "longitude": "120.1"}))

    def test_select_default_sheet_uses_branch_when_present(self) -> None:
        workbook = make_workbook(["branch_name"], [["Branch"]])
        workbook.create_sheet("not_branch")

        with contextlib.redirect_stdout(io.StringIO()):
            selected = cleaner.select_default_sheet(workbook)

        self.assertEqual(selected.title, "branch")

    def test_select_default_sheet_prompts_when_branch_is_absent(self) -> None:
        workbook = Workbook()
        workbook.active.title = "stores"
        workbook.create_sheet("locations")

        with patch("builtins.input", return_value="2"), contextlib.redirect_stdout(io.StringIO()):
            selected = cleaner.select_default_sheet(workbook)

        self.assertEqual(selected.title, "locations")

    def test_prompt_column_selection_returns_selected_column_index(self) -> None:
        workbook = make_workbook(["name", "address_text", "lat_col", "lng_col"], [["Branch", "Address", "", ""]])

        with patch("builtins.input", return_value="3"), contextlib.redirect_stdout(io.StringIO()):
            selected = cleaner.prompt_column_selection(workbook["branch"], "latitude")

        self.assertEqual(selected, 3)

    def test_prompt_actions_supports_both_coordinates_and_address_parsing(self) -> None:
        with patch("builtins.input", return_value="3"), contextlib.redirect_stdout(io.StringIO()):
            actions = cleaner.prompt_actions()

        self.assertEqual(actions, {"coordinates", "parse_address"})

    def test_clean_missing_coordinates_fills_only_invalid_coordinate_rows(self) -> None:
        workbook = make_workbook(
            ["branch_name", "full_address", "latitude", "longitude"],
            [
                ["Valid", "Valid Address", "14.1", "121.1"],
                ["Missing", "Missing Address", "", ""],
                ["No Address", "", "", ""],
                ["No Result", "No Result Address", "", ""],
            ],
        )
        sheet = workbook["branch"]
        geocoder = FakeGeocoder()

        columns = cleaner.CoordinateColumns(latitude=3, longitude=4, address=2)
        summary = cleaner.clean_missing_coordinates(sheet, geocoder, columns=columns)

        self.assertEqual(summary.rows_checked, 4)
        self.assertEqual(summary.rows_needing_coordinates, 3)
        self.assertEqual(summary.geocoded_rows, 1)
        self.assertEqual(summary.skipped_no_address_rows, [4])
        self.assertEqual(summary.no_result_rows, [5])
        self.assertEqual(geocoder.addresses, ["Missing Address", "No Result Address"])
        self.assertEqual(sheet["C2"].value, "14.1")
        self.assertEqual(sheet["D2"].value, "121.1")
        self.assertEqual(sheet["C3"].value, 14.5995)
        self.assertEqual(sheet["D3"].value, 120.9842)

    def test_clean_missing_coordinates_skips_when_all_coordinates_are_valid(self) -> None:
        workbook = make_workbook(
            ["branch_name", "full_address", "latitude", "longitude"],
            [["Valid", "Valid Address", "14.1", "121.1"]],
        )

        summary = cleaner.clean_missing_coordinates(
            workbook["branch"],
            geocoder=None,
            columns=cleaner.CoordinateColumns(latitude=3, longitude=4, address=2),
        )

        self.assertEqual(summary.rows_needing_coordinates, 0)
        self.assertEqual(summary.geocoded_rows, 0)

    def test_clean_missing_coordinates_uses_nonstandard_selected_columns(self) -> None:
        workbook = make_workbook(
            ["name", "address_text", "lat_col", "lng_col", "untouched_lat", "untouched_lng"],
            [["Branch", "Selected Address", "", "", "", ""]],
        )
        sheet = workbook["branch"]
        geocoder = FakeGeocoder()

        summary = cleaner.clean_missing_coordinates(
            sheet,
            geocoder,
            columns=cleaner.CoordinateColumns(latitude=3, longitude=4, address=2),
        )

        self.assertEqual(summary.geocoded_rows, 1)
        self.assertEqual(geocoder.addresses, ["Selected Address"])
        self.assertEqual(sheet["C2"].value, 14.5995)
        self.assertEqual(sheet["D2"].value, 120.9842)
        self.assertEqual(sheet["E2"].value, "")
        self.assertEqual(sheet["F2"].value, "")

    def test_parse_full_addresses_adds_exact_new_columns(self) -> None:
        workbook = make_workbook(
            ["branch_name", "full_address", "latitude", "longitude"],
            [["Branch", "Ropali Plaza, Escriva Drive, Pasig, Metro Manila", "", ""]],
        )
        sheet = workbook["branch"]
        parser = FakeAddressParser()

        summary = cleaner.parse_full_addresses(sheet, parser, address_column=2)
        headers = [cell.value for cell in sheet[1]]

        self.assertEqual(summary.parsed_rows, 1)
        self.assertEqual(parser.addresses, ["Ropali Plaza, Escriva Drive, Pasig, Metro Manila"])
        self.assertEqual(headers[-7:], cleaner.PARSED_ADDRESS_COLUMNS)
        row = {headers[index]: value for index, value in enumerate(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))}
        self.assertEqual(row["barangay2"], "Barangay San Antonio")
        self.assertEqual(row["city_municipality2"], "Pasig")
        self.assertEqual(row["country2"], "Philippines")
        self.assertEqual(row["island_group2"], "Luzon")

    def test_parse_full_addresses_always_appends_fresh_columns_even_when_existing(self) -> None:
        headers = ["branch_name", "full_address", "latitude", "longitude", *cleaner.PARSED_ADDRESS_COLUMNS]
        workbook = make_workbook(headers, [["Branch", "Address", "", "", "", "", "", "", "", ""]])

        cleaner.parse_full_addresses(workbook["branch"], FakeAddressParser(), address_column=2)
        resulting_headers = [cell.value for cell in workbook["branch"][1]]

        self.assertEqual(resulting_headers.count("barangay2"), 2)
        self.assertEqual(resulting_headers[-7:], cleaner.PARSED_ADDRESS_COLUMNS)
        self.assertEqual(workbook["branch"].cell(row=2, column=12).value, "Barangay San Antonio")

    def test_missing_branch_sheet_and_required_columns_raise_clear_errors(self) -> None:
        workbook = Workbook()
        workbook.active.title = "not_branch"

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "dealer.xlsx"
            workbook.save(path)
            with self.assertRaisesRegex(cleaner.BranchCleanerError, "branch"):
                cleaner.load_branch_sheet(path)

        workbook = make_workbook(["branch_name", "latitude", "longitude"], [["Branch", "", ""]])
        with self.assertRaisesRegex(cleaner.BranchCleanerError, "full_address"):
            cleaner.clean_missing_coordinates(workbook["branch"], FakeGeocoder())

        with self.assertRaisesRegex(cleaner.BranchCleanerError, "outside"):
            cleaner.clean_missing_coordinates(
                workbook["branch"],
                FakeGeocoder(),
                columns=cleaner.CoordinateColumns(latitude=20, longitude=3, address=1),
            )

    def test_extract_gemini_json_payload_supports_generate_content_candidates(self) -> None:
        payload = {
            "candidates": [
                {
                    "content": {
                        "parts": [
                            {
                                "text": '{"barangay": null, "city_municipality": "Pasig", "province": "Metro Manila", "region": "NCR", "zip_code": null, "country": "Philippines", "island_group": "Luzon"}'
                            }
                        ]
                    }
                }
            ]
        }

        parsed = cleaner.extract_gemini_json_payload(payload)

        self.assertEqual(parsed["city_municipality"], "Pasig")
        self.assertEqual(parsed["country"], "Philippines")
        self.assertEqual(parsed["island_group"], "Luzon")

    def test_normalize_address_parse_result_canonicalizes_island_group(self) -> None:
        self.assertEqual(
            cleaner.normalize_address_parse_result({"island group": "western visayas"})["island_group"],
            "Visayas",
        )
        self.assertEqual(
            cleaner.normalize_address_parse_result({"island_group": "Mindanao"})["island_group"],
            "Mindanao",
        )
        self.assertIsNone(cleaner.normalize_address_parse_result({"island_group": "NCR"})["island_group"])

    def test_discover_local_gguf_models_reads_models_folder_and_index(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            models_dir = Path(tmp)
            model_dir = models_dir / "test__repo-GGUF"
            model_dir.mkdir(parents=True)
            model_path = model_dir / "model-q4_k_m.gguf"
            model_path.write_bytes(b"fake")
            cleaner.record_model_in_index(
                cleaner.LocalGgufModel("test/repo-GGUF", "model-q4_k_m.gguf", model_path, 4),
                models_dir,
            )

            models = cleaner.discover_local_gguf_models(models_dir)

            self.assertEqual(len(models), 1)
            self.assertEqual(models[0].repo_id, "test/repo-GGUF")
            self.assertEqual(models[0].filename, "model-q4_k_m.gguf")

    def test_search_huggingface_gguf_models_returns_prioritized_gguf_files(self) -> None:
        results = cleaner.search_huggingface_gguf_models("small instruct", api=FakeHfApi())

        self.assertEqual(results[0].repo_id, "test/repo-GGUF")
        self.assertEqual(results[0].filename, "model-q4_k_m.gguf")
        self.assertTrue(all(result.filename.endswith(".gguf") for result in results))

    def test_download_hf_gguf_model_records_model_under_models_dir(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            models_dir = Path(tmp)

            def fake_downloader(repo_id, filename, local_dir, token=None):
                path = Path(local_dir) / filename
                path.write_bytes(b"fake-gguf")
                return str(path)

            local_model = cleaner.download_hf_gguf_model(
                cleaner.HuggingFaceGgufFile("test/repo-GGUF", "model-q4_k_m.gguf", 9),
                models_dir=models_dir,
                downloader=fake_downloader,
            )

            self.assertTrue(local_model.path.exists())
            self.assertEqual(local_model.path.parent.name, "test__repo-GGUF")
            self.assertTrue((models_dir / "models_index.json").exists())

    def test_llama_cpp_address_parser_extracts_json_and_retries_once(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            model_path = Path(tmp) / "fake.gguf"
            model_path.write_bytes(b"fake")
            fake_llama = FakeLlama(
                [
                    "not json",
                    '{"barangay": null, "city_municipality": "Pasig", "province": "Metro Manila", "region": "NCR", "zip_code": null, "country": "Philippines", "island_group": "Luzon"}',
                ]
            )

            parser = cleaner.LlamaCppAddressParser(model_path, llama_factory=lambda **kwargs: fake_llama)
            parsed = parser.parse_address("Ropali Plaza, Pasig")

            self.assertEqual(parsed["city_municipality"], "Pasig")
            self.assertEqual(len(fake_llama.prompts), 2)

    def test_llama_cpp_address_parser_uses_chat_completion_when_available(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            model_path = Path(tmp) / "fake.gguf"
            model_path.write_bytes(b"fake")
            fake_llama = FakeChatLlama(
                [
                    '```json\n{"barangay": null, "city_municipality": "Pasig", "province": "Metro Manila", "region": "NCR", "zip_code": null, "country": "Philippines", "island_group": "Luzon"}\n```'
                ]
            )

            parser = cleaner.LlamaCppAddressParser(model_path, llama_factory=lambda **kwargs: fake_llama)
            parsed = parser.parse_address("Ropali Plaza, Escriva Drive, Pasig, Metro Manila")

            self.assertEqual(parsed["city_municipality"], "Pasig")
            self.assertIn("strict JSON", fake_llama.messages[0][0]["content"])

    def test_build_parser_from_prompt_can_choose_local_provider(self) -> None:
        with (
            patch("builtins.input", return_value="local"),
            patch("branch_cleaner.branch_sheet_cleaner.build_local_parser_from_prompt", return_value=FakeAddressParser()) as local_builder,
        ):
            parser = cleaner.build_parser_from_prompt()

        self.assertIsInstance(parser, FakeAddressParser)
        local_builder.assert_called_once()

    def test_build_local_parser_from_prompt_uses_isolated_worker(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            model_path = Path(tmp) / "fake.gguf"
            model_path.write_bytes(b"fake")
            local_model = cleaner.LocalGgufModel("models", "fake.gguf", model_path, 4)

            with (
                patch("branch_cleaner.branch_sheet_cleaner.prompt_local_model", return_value=local_model),
                patch("builtins.input", side_effect=["", ""]),
                contextlib.redirect_stdout(io.StringIO()),
            ):
                parser = cleaner.build_local_parser_from_prompt(Path(tmp))

            self.assertIsInstance(parser, cleaner.LlamaCppWorkerAddressParser)
            self.assertEqual(parser.n_gpu_layers, 0)
            self.assertFalse(parser.flash_attn)

    def test_prompt_local_runtime_allows_gpu_and_flash_choices(self) -> None:
        with patch("builtins.input", side_effect=["-1", "y"]), contextlib.redirect_stdout(io.StringIO()):
            gpu_layers, flash_attn = cleaner.prompt_local_runtime()

        self.assertEqual(gpu_layers, -1)
        self.assertTrue(flash_attn)

    def test_llama_cpp_worker_parser_extracts_worker_json(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            model_path = Path(tmp) / "fake.gguf"
            model_path.write_bytes(b"fake")
            completed = cleaner.subprocess.CompletedProcess(
                args=[],
                returncode=0,
                stdout=json_payload("Pasig"),
                stderr="",
            )

            with patch("branch_cleaner.branch_sheet_cleaner.subprocess.run", return_value=completed) as run:
                parser = cleaner.LlamaCppWorkerAddressParser(model_path, python_executable="python-test", persistent=False)
                parsed = parser.parse_address("Ropali Plaza, Pasig")

            self.assertEqual(parsed["city_municipality"], "Pasig")
            self.assertIn("--local-parse-worker", run.call_args.args[0])

    def test_llama_cpp_worker_parser_reports_crashed_worker(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            model_path = Path(tmp) / "fake.gguf"
            model_path.write_bytes(b"fake")
            completed = cleaner.subprocess.CompletedProcess(
                args=[],
                returncode=139,
                stdout="",
                stderr="segmentation fault",
            )

            with patch("branch_cleaner.branch_sheet_cleaner.subprocess.run", return_value=completed):
                parser = cleaner.LlamaCppWorkerAddressParser(model_path, persistent=False)
                with self.assertRaisesRegex(cleaner.BranchCleanerError, "exit code 139"):
                    parser.parse_address("Ropali Plaza, Pasig")

    def test_llama_cpp_worker_parser_uses_persistent_worker(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            model_path = Path(tmp) / "fake.gguf"
            model_path.write_bytes(b"fake")

            class FakeStdout:
                def __init__(self):
                    self.lines = [
                        '{"ok": true, "ready": true}\n',
                        json_payload("Pasig") + "\n",
                    ]

                def readline(self):
                    return self.lines.pop(0)

                def fileno(self):
                    return 0

            class FakeStdin:
                def __init__(self):
                    self.writes: list[str] = []

                def write(self, value):
                    self.writes.append(value)

                def flush(self):
                    return None

            class FakeProcess:
                def __init__(self):
                    self.stdin = FakeStdin()
                    self.stdout = FakeStdout()

                def poll(self):
                    return None

                def terminate(self):
                    return None

                def wait(self, timeout=None):
                    return 0

            class FakeSelector:
                def register(self, *args, **kwargs):
                    return None

                def select(self, timeout=None):
                    return [object()]

                def close(self):
                    return None

            fake_process = FakeProcess()
            with (
                patch("branch_cleaner.branch_sheet_cleaner.subprocess.Popen", return_value=fake_process),
                patch("branch_cleaner.branch_sheet_cleaner.selectors.DefaultSelector", return_value=FakeSelector()),
            ):
                parser = cleaner.LlamaCppWorkerAddressParser(model_path)
                parsed = parser.parse_address("Ropali Plaza, Pasig")
                parser.close()

            self.assertEqual(parsed["city_municipality"], "Pasig")
            self.assertIn("model_path", fake_process.stdin.writes[0])
            self.assertIn("Ropali Plaza", fake_process.stdin.writes[1])

    def test_csv_file_can_be_loaded_cleaned_and_saved(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "branches.csv"
            output = Path(tmp) / "branches_cleaned.csv"
            with source.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["name", "address_text", "lat_col", "lng_col"])
                writer.writerow(["Branch", "Selected Address", "", ""])

            loaded = cleaner.load_file_for_cleaning(source)
            cleaner.clean_missing_coordinates(
                loaded.sheet,
                FakeGeocoder(),
                columns=cleaner.CoordinateColumns(latitude=3, longitude=4, address=2),
            )
            cleaner.parse_full_addresses(loaded.sheet, FakeAddressParser(), address_column=2)
            cleaner.save_loaded_file(loaded, output)

            with output.open("r", newline="", encoding="utf-8") as handle:
                rows = list(csv.reader(handle))

            self.assertEqual(rows[0][-7:], cleaner.PARSED_ADDRESS_COLUMNS)
            self.assertEqual(rows[1][2], "14.5995")
            self.assertEqual(rows[1][4], "Barangay San Antonio")

    def test_saving_cleaned_copy_preserves_non_branch_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "dealer.xlsx"
            output = Path(tmp) / "dealer_cleaned.xlsx"
            workbook = make_workbook(
                ["branch_name", "full_address", "latitude", "longitude"],
                [["Branch", "Address", "", ""]],
            )
            workbook.save(source)

            loaded, sheet = cleaner.load_branch_sheet(source)
            cleaner.clean_missing_coordinates(sheet, FakeGeocoder(), columns=cleaner.CoordinateColumns(latitude=3, longitude=4, address=2))
            loaded.save(output)

            original = load_workbook(source, read_only=True, data_only=True)
            cleaned = load_workbook(output, read_only=True, data_only=True)
            self.assertIn("other_sheet", cleaned.sheetnames)
            self.assertEqual(original["branch"]["C2"].value, None)
            self.assertEqual(cleaned["branch"]["C2"].value, 14.5995)
            self.assertEqual(cleaned["other_sheet"]["A2"].value, "yes")


if __name__ == "__main__":
    unittest.main()
